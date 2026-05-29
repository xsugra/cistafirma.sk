from io import BytesIO

from django.http import HttpResponse
from django.utils.html import format_html


class AdminDisplayMixin:
    @staticmethod
    def badge(label: str, variant: str) -> str:
        return format_html(
            '<span class="cf-badge cf-badge--{}">'
            '<span class="cf-badge__dot"></span>{}</span>',
            variant, label,
        )

    @staticmethod
    def progress_bar(percentage: float, fill_variant: str = "blue") -> str:
        clamped = min(percentage, 100)
        return format_html(
            '<div class="cf-progress">'
            '<div class="cf-progress__fill cf-progress__fill--{}" style="width:{}%">'
            '{}%</div></div>',
            fill_variant, clamped, percentage,
        )


class XlsxExportMixin:
    EXPORT_FIELDS: list[tuple[str, str]] = []

    def _get_export_data(self, queryset):
        data = []
        for obj in queryset:
            row = []
            for field_name, _ in self.EXPORT_FIELDS:
                value = getattr(obj, field_name, "")
                if value is None:
                    value = ""
                elif isinstance(value, bool):
                    value = "Ano" if value else "Nie"
                row.append(value)
            data.append(row)
        return data

    def _build_xlsx_response(self, queryset, filename: str) -> HttpResponse:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return HttpResponse("Chyba kniznica openpyxl.", status=500)

        wb = Workbook()
        ws = wb.active
        ws.title = "Export"

        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")

        headers = [label for _, label in self.EXPORT_FIELDS]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill

        for row_idx, row_data in enumerate(self._get_export_data(queryset), 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(value) if value else "")

        for col in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)

        output = BytesIO()
        wb.save(output)
        output.seek(0)

        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        return response
