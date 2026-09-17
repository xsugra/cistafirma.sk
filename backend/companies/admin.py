from django.contrib import admin
from django.core.cache import cache
from django.utils.html import format_html
from django.utils.safestring import mark_safe
from django.utils import timezone
from django.contrib import messages
from django.db.models import Count, Exists, OuterRef, Q
from django.http import HttpResponse, HttpResponseRedirect, StreamingHttpResponse
from django.urls import reverse
import csv
from io import BytesIO
from .models import (
    Company,
    CompanyFinancialResult,
    PostalCodeArea,
    SectorBenchmark,
    LEGAL_FORMS_SHORT,
    LEGAL_FORMS_CHOICES,
    LEGAL_FORMS,
)
from registers.models import OrsrCompanyProfile
from core.formatting import format_currency_eur, format_int_space

MAX_SYNC_EXPORT_ROWS = 5_000


class _CSVBuffer:
    """Minimal csv.writer target that returns each generated row for streaming."""

    def write(self, value):
        return value


try:
    from unfold.admin import ModelAdmin as UnfoldModelAdmin
    from unfold.decorators import action as unfold_action
    UNFOLD_AVAILABLE = True
except ImportError:
    UnfoldModelAdmin = admin.ModelAdmin
    unfold_action = admin.action
    UNFOLD_AVAILABLE = False


class LegalFormFilter(admin.SimpleListFilter):
    title = 'Pravna forma'
    parameter_name = 'pravna_forma'

    def lookups(self, request, model_admin):
        try:
            used_forms_queryset = Company.objects.exclude(
                pravna_forma__isnull=True
            ).exclude(
                pravna_forma__exact=''
            ).values_list('pravna_forma', flat=True).distinct().order_by('pravna_forma')
            used_forms = sorted(set(used_forms_queryset))
            lookups = []
            for form_code in used_forms:
                form_code_str = str(form_code).strip()
                if form_code_str:
                    form_short = LEGAL_FORMS_SHORT.get(form_code_str, f'Neznama ({form_code_str})')
                    lookups.append((form_code_str, f'{form_code_str} - {form_short}'))
            return lookups
        except Exception:
            return [
                ('112', '112 - s. r. o.'),
                ('121', '121 - a. s.'),
                ('101', '101 - FO-podnikatel'),
                ('111', '111 - v. o. s.'),
            ]

    def queryset(self, request, queryset):
        if self.value():
            return queryset.filter(pravna_forma=self.value())
        return queryset


class RokZalozeniaFilter(admin.SimpleListFilter):
    title = 'Rok zalozenia'
    parameter_name = 'rok_zalozenia'

    def lookups(self, request, model_admin):
        try:
            years_queryset = Company.objects.exclude(
                datum_zalozenia__isnull=True
            ).dates('datum_zalozenia', 'year', order='DESC')
            return [(str(d.year), str(d.year)) for d in years_queryset]
        except Exception:
            import datetime
            current_year = datetime.date.today().year
            return [(str(y), str(y)) for y in range(current_year, current_year - 30, -1)]

    def queryset(self, request, queryset):
        if self.value():
            return queryset.filter(datum_zalozenia__year=int(self.value()))
        return queryset


class HasDebtFilter(admin.SimpleListFilter):
    title = 'Stav dlhov'
    parameter_name = 'has_debt'

    def lookups(self, request, model_admin):
        return [
            ('yes', 'S dlhmi'),
            ('no', 'Bez dlhov'),
        ]

    def queryset(self, request, queryset):
        from django.db.models import Q
        if self.value() == 'yes':
            return queryset.filter(
                Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0)
            )
        if self.value() == 'no':
            return queryset.exclude(
                Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0)
            )
        return queryset


class DataCompletenessFilter(admin.SimpleListFilter):
    title = 'Kompletnost dat'
    parameter_name = 'data_completeness'

    def lookups(self, request, model_admin):
        return [
            ('has_orsr', 'S ORSR profilom'),
            ('no_orsr', 'Bez ORSR profilu'),
            ('has_financials', 'S financnymi vysledkami'),
            ('no_financials', 'Bez financnych vysledkov'),
        ]

    def queryset(self, request, queryset):
        if self.value() == 'has_orsr':
            return queryset.filter(orsr_profile__isnull=False)
        if self.value() == 'no_orsr':
            return queryset.filter(orsr_profile__isnull=True)
        if self.value() == 'has_financials':
            return queryset.filter(financial_results__isnull=False).distinct()
        if self.value() == 'no_financials':
            return queryset.filter(financial_results__isnull=True)
        return queryset


class OrsrCompanyProfileInline(admin.StackedInline):
    model = OrsrCompanyProfile
    can_delete = False
    extra = 0
    max_num = 1
    readonly_fields = [
        'ico', 'oddiel', 'vlozka_cislo', 'obchodne_meno', 'sidlo',
        'den_zapisu', 'pravna_forma', 'konanie_menom_spolocnosti',
        'vyska_zakladneho_imania', 'predmet_podnikania', 'spolocnici',
        'vklady_spolocnikov', 'statutarny_organ', 'prokura', 'orsr_aktualizacia_dat',
        'orsr_datum_vypisu', 'source_url', 'fetch_ok', 'last_error',
        'last_synced_at', 'raw_sections', 'raw_payload',
    ]
    fieldsets = (
        ('ORSR zaklad', {
            'fields': (
                ('ico', 'oddiel', 'vlozka_cislo'),
                'obchodne_meno',
                'sidlo',
                ('den_zapisu', 'pravna_forma'),
            ),
        }),
        ('ORSR osoby a vazby', {
            'fields': ('statutarny_organ', 'prokura', 'spolocnici', 'vklady_spolocnikov'),
            'classes': ('collapse',),
        }),
        ('ORSR podnikanie a kapital', {
            'fields': ('predmet_podnikania', 'konanie_menom_spolocnosti', 'vyska_zakladneho_imania'),
            'classes': ('collapse',),
        }),
        ('ORSR meta', {
            'fields': (
                ('orsr_aktualizacia_dat', 'orsr_datum_vypisu'),
                'source_url',
                ('fetch_ok', 'last_synced_at'),
                'last_error',
                'raw_sections',
                'raw_payload',
            ),
            'classes': ('collapse',),
        }),
    )


class CompanyFinancialResultInline(admin.TabularInline):
    model = CompanyFinancialResult
    extra = 0
    # Both P&L result rows: `profit` is the operating result and
    # `profit_after_tax` is the bottom line, and an inline that shows one of
    # them invites the reader to take it for the other.
    fields = ('year', 'revenue', 'profit', 'profit_after_tax', 'source', 'updated_at')
    readonly_fields = ('updated_at',)
    ordering = ('-year',)


@admin.register(Company)
class CompanyAdmin(UnfoldModelAdmin):
    list_display = [
        'ico', 'nazov_UJ', 'mesto', 'legal_form_display',
        'datum_zalozenia', 'status_display',
        'vat_payer_display', 'tax_reliability_display',
        'risk_display', 'data_quality_display',
        'datum_poslednej_upravy',
    ]

    list_filter = (LegalFormFilter, RokZalozeniaFilter, HasDebtFilter, DataCompletenessFilter)

    list_display_links = ['ico', 'nazov_UJ']
    inlines = [OrsrCompanyProfileInline, CompanyFinancialResultInline]
    change_list_template = 'admin/companies/company/change_list.html'
    change_form_template = 'admin/companies/company/change_form.html'

    search_fields = [
        'ico', 'nazov_UJ', 'dic', 'ic_dph', 'mesto', 'ulica', 'ruz_id',
    ]
    search_help_text = "Hladanie: ICO, nazov, DIC, IC DPH, mesto, ulica, RUZ ID"


    ordering = ['-datum_poslednej_upravy', 'nazov_UJ']
    list_per_page = 50
    list_max_show_all = 500

    readonly_fields = [
        'ruz_id', 'datum_poslednej_upravy', 'last_insurance_debt', 'fs_update_date',
        'id_uctovnych_zavierok', 'id_vyrocnych_sprav',
    ]

    fieldsets = (
        ('Zakladne udaje', {
            'fields': (
                ('ico', 'dic', 'sid'),
                'nazov_UJ',
                'ruz_id',
            )
        }),
        ('Adresa', {
            'fields': (
                ('ulica', 'mesto'),
                ('psc', 'okres', 'kraj'),
                'sidlo',
            )
        }),
        ('Klasifikacia', {
            'fields': (
                ('pravna_forma', 'sk_NACE'),
                ('velkost_organizacie', 'druh_vlastnictva'),
                'konsolidovana',
            )
        }),
        ('Datumy', {
            'fields': (
                ('datum_zalozenia', 'datum_zrusenia'),
                'datum_poslednej_upravy',
            )
        }),
        ('DPH a Financna sprava', {
            'fields': (
                ('vat_payer', 'ic_dph'),
                ('datum_reg_dph', 'vat_deleted_date'),
                'vat_deleted_reason',
                'tax_reliability',
                'tax_debt',
                'bank_accounts',
                'fs_update_date',
            ),
            'classes': ('collapse',),
        }),
        ('Dlhy v poistovniach', {
            'fields': (
                ('debt_vszp', 'debt_soc_poist'),
                'social_listed_without_amount',
                'last_insurance_debt',
            ),
            'classes': ('collapse',),
        }),
        ('Uctovne zavierky a Vyrocne spravy', {
            'fields': (
                'id_uctovnych_zavierok',
                'id_vyrocnych_sprav',
            ),
            'classes': ('collapse',),
        }),
        ('Zdroj dat', {
            'fields': ('zdroj_dat',),
            'classes': ('collapse',),
        }),
    )

    # ── Display Methods ──

    @admin.display(description='Stav')
    def status_display(self, obj):
        if obj.datum_zrusenia:
            return mark_safe(
                '<span class="cf-badge cf-badge--danger">'
                '<span class="cf-badge__dot"></span>Zrusena</span>'
            )
        return mark_safe(
            '<span class="cf-badge cf-badge--success">'
            '<span class="cf-badge__dot"></span>Aktivna</span>'
        )

    @admin.display(description='DPH', boolean=True)
    def vat_payer_display(self, obj):
        return obj.vat_payer

    @admin.display(description='Pravna forma')
    def legal_form_display(self, obj):
        if not obj.pravna_forma:
            return '-'
        form_short = LEGAL_FORMS_SHORT.get(str(obj.pravna_forma), '?')
        return format_html(
            '<span title="{}">{}</span>',
            f'{obj.pravna_forma}', form_short
        )

    @admin.display(description='Dan. spolahlivost')
    def tax_reliability_display(self, obj):
        if not obj.tax_reliability:
            return mark_safe('<span class="cf-risk cf-risk--unknown">-</span>')
        mapping = {
            'vysoko spoľahlivý': ('cf-badge--success', 'Vysoko'),
            'spoľahlivý': ('cf-badge--info', 'OK'),
            'nespoľahlivý': ('cf-badge--danger', 'Nespolahlivy'),
        }
        css, label = mapping.get(obj.tax_reliability.lower(), ('cf-badge--idle', obj.tax_reliability))
        return format_html('<span class="cf-badge {}">{}</span>', css, label)

    @admin.display(description='Riziko')
    def risk_display(self, obj):
        issues = []
        total_debt = 0
        if obj.debt_vszp and obj.debt_vszp > 0:
            issues.append(f'VSZP {format_currency_eur(obj.debt_vszp)}')
            total_debt += obj.debt_vszp
        if obj.debt_soc_poist and obj.debt_soc_poist > 0:
            issues.append(f'SP {format_currency_eur(obj.debt_soc_poist)}')
            total_debt += obj.debt_soc_poist
        if obj.tax_debt and obj.tax_debt > 0:
            issues.append(f'Dan {format_currency_eur(obj.tax_debt)}')
            total_debt += obj.tax_debt

        if issues:
            level = 'cf-badge--danger' if total_debt > 1000 else 'cf-badge--warning'
            return format_html(
                '<span class="{}" title="{}">'
                '<span class="cf-badge__dot"></span>{}</span>',
                f'cf-badge {level}',
                ', '.join(issues).replace('&euro;', '€'),
                format_currency_eur(total_debt)
            )
        if obj.social_listed_without_amount:
            # Listed by the social insurer, with no sum published for it. The
            # badge below is shaped like money, so this is not folded into it:
            # "0,00 €" beside a listing that is not about money would be a new
            # wrong answer in place of the old one.
            return mark_safe(
                '<span class="cf-badge cf-badge--warning" title="Sociálna '
                'poisťovňa uvádza spoločnosť v zozname dlžníkov bez zverejnenej '
                'sumy (nesplnená vykazovacia povinnosť)">SP bez sumy</span>'
            )
        if obj.last_insurance_debt:
            return mark_safe(
                '<span class="cf-badge cf-badge--success">OK</span>'
            )
        return mark_safe('<span class="cf-risk cf-risk--unknown">-</span>')

    @admin.display(description='Data')
    def data_quality_display(self, obj):
        parts = []
        has_orsr = bool(getattr(obj, '_has_orsr', False))
        has_financials = bool(getattr(obj, '_has_financials', False))

        if has_orsr:
            parts.append('<span title="ORSR profil" style="color:var(--cf-emerald-500)">OR</span>')
        else:
            parts.append('<span title="Chyba ORSR" style="color:var(--cf-slate-300)">OR</span>')

        if has_financials:
            parts.append('<span title="Financne vysledky" style="color:var(--cf-emerald-500)">FIN</span>')
        else:
            parts.append('<span title="Chybaju financie" style="color:var(--cf-slate-300)">FIN</span>')

        if obj.vat_payer is not None:
            parts.append('<span title="DPH info" style="color:var(--cf-emerald-500)">DPH</span>')
        else:
            parts.append('<span title="Chyba DPH" style="color:var(--cf-slate-300)">DPH</span>')

        return format_html(
            '<span style="font-size:11px;font-weight:600;display:flex;gap:4px;">{}</span>',
            mark_safe(' '.join(parts))
        )

    # ── Actions ──

    actions = [
        'sync_from_ruz',
        'sync_from_orsr',
        'sync_financials_from_ruz',
        'check_insurance_debts',
        'check_fs_data',
        'refresh_all_data',
        'export_to_csv',
        'export_to_xlsx',
        'export_filtered_to_csv',
        'export_filtered_to_xlsx',
    ]

    @admin.action(description='Sync z RUZ API')
    def sync_from_ruz(self, request, queryset):
        from registers.tasks import sync_single_company_from_ruz
        count = 0
        for company in queryset:
            sync_single_company_from_ruz.delay(company.ico)
            count += 1
        self.message_user(request, f'Naplanovana synchronizacia z RUZ pre {count} firiem.', messages.SUCCESS)

    @admin.action(description='Sync z ORSR')
    def sync_from_orsr(self, request, queryset):
        from registers.tasks import sync_company_orsr_data
        count = 0
        for company in queryset:
            sync_company_orsr_data.delay(company.id)
            count += 1
        self.message_user(request, f'Naplanovana synchronizacia z ORSR pre {count} firiem.', messages.SUCCESS)

    @admin.action(description='Sync hosp. vysledky z RUZ')
    def sync_financials_from_ruz(self, request, queryset):
        from registers.tasks import sync_company_financials_from_ruz
        count = 0
        for company in queryset:
            sync_company_financials_from_ruz.delay(company.id)
            count += 1
        self.message_user(request, f'Naplanovana synchronizacia hosp. vysledkov pre {count} firiem.', messages.SUCCESS)

    @admin.action(description='Kontrola dlhov v poistovniach')
    def check_insurance_debts(self, request, queryset):
        from registers.tasks import update_insurance_debt
        count = 0
        for company in queryset:
            update_insurance_debt.delay(company.id)
            count += 1
        self.message_user(request, f'Naplanovana kontrola dlhov pre {count} firiem.', messages.SUCCESS)

    @admin.action(description='Aktualizovat z Financnej spravy')
    def check_fs_data(self, request, queryset):
        from registers.tasks import update_fs_data_task
        update_fs_data_task.delay()
        self.message_user(request, 'Spustena aktualizacia z Financnej spravy.', messages.SUCCESS)

    @admin.action(description='Kompletny refresh dat')
    def refresh_all_data(self, request, queryset):
        from registers.tasks import (
            sync_single_company_from_ruz,
            sync_company_financials_from_ruz,
            sync_company_orsr_data,
            update_insurance_debt,
        )
        count = 0
        for company in queryset:
            sync_single_company_from_ruz.delay(company.ico)
            sync_company_orsr_data.delay(company.id)
            sync_company_financials_from_ruz.delay(company.id)
            update_insurance_debt.delay(company.id)
            count += 1
        self.message_user(request, f'Naplanovany kompletny refresh pre {count} firiem.', messages.SUCCESS)

    # ── Export Fields ──

    EXPORT_FIELDS = [
        ('ico', 'ICO'),
        ('nazov_UJ', 'Nazov'),
        ('dic', 'DIC'),
        ('ic_dph', 'IC DPH'),
        ('ulica', 'Ulica'),
        ('mesto', 'Mesto'),
        ('psc', 'PSC'),
        ('okres', 'Okres'),
        ('kraj', 'Kraj'),
        ('pravna_forma', 'Pravna forma'),
        ('sk_NACE', 'SK NACE'),
        ('velkost_organizacie', 'Velkost organizacie'),
        ('datum_zalozenia', 'Datum zalozenia'),
        ('datum_zrusenia', 'Datum zrusenia'),
        ('vat_payer', 'Platitel DPH'),
        ('tax_reliability', 'Danova spolahlivost'),
        ('tax_debt', 'Danovy dlh'),
        ('debt_vszp', 'Dlh VSZP'),
        ('debt_soc_poist', 'Dlh Socialna poistovna'),
        ('datum_poslednej_upravy', 'Posledna aktualizacia'),
    ]

    def _get_export_data(self, queryset):
        for company in queryset[:MAX_SYNC_EXPORT_ROWS].iterator(chunk_size=500):
            row = []
            for field_name, _ in self.EXPORT_FIELDS:
                value = getattr(company, field_name, '')
                if value is None:
                    value = ''
                elif isinstance(value, bool):
                    value = 'Ano' if value else 'Nie'
                row.append(value)
            yield row

    def _export_count(self, request, queryset):
        count = queryset.count()
        if count <= MAX_SYNC_EXPORT_ROWS:
            return count
        self.message_user(
            request,
            f'Export je obmedzený na {MAX_SYNC_EXPORT_ROWS} riadkov. Zúžte filtre a skúste znova.',
            messages.ERROR,
        )
        return None

    def get_queryset(self, request):
        queryset = super().get_queryset(request).annotate(
            _has_orsr=Exists(OrsrCompanyProfile.objects.filter(company_id=OuterRef('pk'))),
            _has_financials=Exists(CompanyFinancialResult.objects.filter(company_id=OuterRef('pk'))),
        )
        return self.get_filtered_queryset(request, queryset)

    def get_filtered_queryset(self, request, queryset=None):
        queryset = queryset or Company.objects.all()
        search_query = request.GET.get('q', '')
        if search_query:
            from django.db.models import Q
            queryset = queryset.filter(
                Q(ico__icontains=search_query) |
                Q(nazov_UJ__icontains=search_query) |
                Q(dic__icontains=search_query) |
                Q(ic_dph__icontains=search_query) |
                Q(mesto__icontains=search_query) |
                Q(ulica__icontains=search_query) |
                Q(ruz_id__icontains=search_query)
            )
        text_filters = {
            'mesto': 'mesto__icontains',
            'psc': 'psc__icontains',
            'sk_NACE': 'sk_NACE__icontains',
            'sidlo': 'sidlo__icontains',
            'okres': 'okres__icontains',
        }
        exact_filters = {
            'pravna_forma': 'pravna_forma',
            'kraj': 'kraj',
            'velkost_organizacie': 'velkost_organizacie',
            'vat_payer': 'vat_payer',
            'tax_reliability': 'tax_reliability',
            'konsolidovana': 'konsolidovana',
        }
        for param, field in text_filters.items():
            value = request.GET.get(param)
            if value:
                queryset = queryset.filter(**{field: value})
        for param, field in exact_filters.items():
            value = request.GET.get(param)
            if value and value != 'all':
                if field == 'vat_payer':
                    if value in ('1', 'true', 'True'):
                        queryset = queryset.filter(vat_payer=True)
                    elif value in ('0', 'false', 'False'):
                        queryset = queryset.filter(vat_payer=False)
                    else:
                        queryset = queryset.filter(vat_payer=value)
                elif field == 'konsolidovana':
                    if value in ('1', 'true', 'True'):
                        queryset = queryset.filter(konsolidovana=True)
                    elif value in ('0', 'false', 'False'):
                        queryset = queryset.filter(konsolidovana=False)
                    else:
                        queryset = queryset.filter(konsolidovana=value)
                else:
                    queryset = queryset.filter(**{field: value})
        active = request.GET.get('active')
        if active == '1':
            queryset = queryset.filter(datum_zrusenia__isnull=True)
        elif active == '0':
            queryset = queryset.filter(datum_zrusenia__isnull=False)

        has_debt = request.GET.get('has_debt')
        if has_debt == 'yes':
            queryset = queryset.filter(
                Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0)
            )
        elif has_debt == 'no':
            queryset = queryset.exclude(
                Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0)
            )

        has_orsr = request.GET.get('has_orsr')
        if has_orsr == '1':
            queryset = queryset.filter(orsr_profile__isnull=False)
        elif has_orsr == '0':
            queryset = queryset.filter(orsr_profile__isnull=True)

        has_financials = request.GET.get('has_financials')
        if has_financials == '1':
            queryset = queryset.filter(financial_results__isnull=False).distinct()
        elif has_financials == '0':
            queryset = queryset.filter(financial_results__isnull=True)

        legal_form = request.GET.get('legal_form')
        if legal_form and legal_form != 'all':
            queryset = queryset.filter(pravna_forma=legal_form)

        founded_year = request.GET.get('founded_year')
        if founded_year and founded_year != 'all':
            queryset = queryset.filter(datum_zalozenia__year=int(founded_year))

        debt_state = request.GET.get('debt_state')
        if debt_state == 'debt_free':
            queryset = queryset.exclude(Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0))
        elif debt_state == 'has_debt':
            queryset = queryset.filter(Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0))

        data_state = request.GET.get('data_state')
        if data_state == 'complete':
            queryset = queryset.filter(orsr_profile__isnull=False, financial_results__isnull=False).distinct()
        elif data_state == 'missing_orsr':
            queryset = queryset.filter(orsr_profile__isnull=True)
        elif data_state == 'missing_financials':
            queryset = queryset.filter(financial_results__isnull=True)

        debt_empty = request.GET.get('debt_empty')
        if debt_empty == '1':
            queryset = queryset.filter(debt_vszp__isnull=True, debt_soc_poist__isnull=True, tax_debt__isnull=True)
        elif debt_empty == '0':
            queryset = queryset.exclude(debt_vszp__isnull=True, debt_soc_poist__isnull=True, tax_debt__isnull=True)

        isnull_mappings = {
            'datum_zrusenia__isempty': 'datum_zrusenia__isnull',
            'debt_vszp__isempty': 'debt_vszp__isnull',
            'debt_soc_poist__isempty': 'debt_soc_poist__isnull',
            'tax_debt__isempty': 'tax_debt__isnull',
        }
        for param, filter_field in isnull_mappings.items():
            value = request.GET.get(param)
            if value == '1':
                queryset = queryset.filter(**{filter_field: True})
            elif value == '0':
                queryset = queryset.filter(**{filter_field: False})
        return queryset

    @admin.action(description='Export vybranych do CSV')
    def export_to_csv(self, request, queryset):
        count = self._export_count(request, queryset)
        if count is None:
            return None
        response = self._export_queryset_csv(queryset)
        self.message_user(request, f'Exportovanych {count} firiem do CSV.', messages.SUCCESS)
        return response

    @admin.action(description='Export vsetkych filtrovanych do CSV')
    def export_filtered_to_csv(self, request, queryset):
        filtered_queryset = self.get_filtered_queryset(request)
        count = self._export_count(request, filtered_queryset)
        if count is None:
            return None
        response = self._export_queryset_csv(filtered_queryset)
        self.message_user(request, f'Exportovanych {count} filtrovanych firiem do CSV.', messages.SUCCESS)
        return response

    @admin.action(description='Export vybranych do XLSX')
    def export_to_xlsx(self, request, queryset):
        count = self._export_count(request, queryset)
        if count is None:
            return None
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            self.message_user(request, 'Chyba kniznica openpyxl. Nainstalujte: pip install openpyxl', messages.ERROR)
            return
        wb = Workbook()
        ws = wb.active
        ws.title = 'Firmy'
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='DAEEF3', end_color='DAEEF3', fill_type='solid')
        headers = [label for _, label in self.EXPORT_FIELDS]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        for row_idx, row_data in enumerate(self._get_export_data(queryset), 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="firmy_export.xlsx"'
        self.message_user(request, f'Exportovanych {count} firiem do XLSX.', messages.SUCCESS)
        return response

    @admin.action(description='Export vsetkych filtrovanych do XLSX')
    def export_filtered_to_xlsx(self, request, queryset):
        filtered_queryset = self.get_filtered_queryset(request)
        count = self._export_count(request, filtered_queryset)
        if count is None:
            return None
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            self.message_user(request, 'Chyba kniznica openpyxl. Nainstalujte: pip install openpyxl', messages.ERROR)
            return
        wb = Workbook()
        ws = wb.active
        ws.title = 'Firmy'
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='DAEEF3', end_color='DAEEF3', fill_type='solid')
        headers = [label for _, label in self.EXPORT_FIELDS]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        for row_idx, row_data in enumerate(self._get_export_data(filtered_queryset), 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="firmy_filtrovane_export.xlsx"'
        self.message_user(request, f'Exportovanych {count} filtrovanych firiem do XLSX.', messages.SUCCESS)
        return response

    # ── Custom Admin Views ──

    def get_urls(self):
        from django.urls import path
        urls = super().get_urls()
        custom_urls = [
            path(
                '<int:company_id>/sync-now/',
                self.admin_site.admin_view(self.sync_now_view),
                name='companies_company_sync_now',
            ),
            path(
                'add-from-ruz/',
                self.admin_site.admin_view(self.add_company_from_ruz_view),
                name='companies_company_add_from_ruz',
            ),
            path(
                'trigger-full-ruz-sync/',
                self.admin_site.admin_view(self.trigger_full_ruz_sync_view),
                name='companies_company_trigger_full_ruz_sync',
            ),
            path(
                'export-filtered/',
                self.admin_site.admin_view(self.export_filtered_view),
                name='companies_company_export_filtered',
            ),
        ]
        return custom_urls + urls

    def sync_now_view(self, request, company_id: int):
        from registers.tasks import sync_company_now
        try:
            company = Company.objects.get(id=company_id)
        except Company.DoesNotExist:
            self.message_user(request, 'Firma neexistuje.', messages.ERROR)
            return HttpResponseRedirect(reverse('admin:companies_company_changelist'))
        try:
            sync_company_now.delay(company.id)
            self.message_user(
                request,
                f'Full sync naplanovany pre {company.nazov_UJ} ({company.ico}).',
                messages.SUCCESS,
            )
        except Exception as exc:
            self.message_user(
                request,
                f'Nepodarilo sa spustit sync: {str(exc)[:120]}',
                messages.ERROR,
            )
        return HttpResponseRedirect(reverse('admin:companies_company_change', args=[company.id]))

    def export_filtered_view(self, request):
        export_format = request.GET.get('format', 'csv')
        queryset = self.get_filtered_queryset(request)
        count = self._export_count(request, queryset)
        if count is None:
            return HttpResponse(
                f'Export je obmedzený na {MAX_SYNC_EXPORT_ROWS} riadkov. Zúžte filtre a skúste znova.',
                status=400,
                content_type='text/plain; charset=utf-8',
            )
        if export_format == 'xlsx':
            return self._export_queryset_xlsx(queryset)
        return self._export_queryset_csv(queryset)

    def _export_queryset_csv(self, queryset):
        def stream_rows():
            writer = csv.writer(_CSVBuffer(), delimiter=';')
            yield '\ufeff'
            yield writer.writerow([label for _, label in self.EXPORT_FIELDS])
            for row in self._get_export_data(queryset):
                yield writer.writerow(row)

        response = StreamingHttpResponse(stream_rows(), content_type='text/csv; charset=utf-8')
        response['Content-Disposition'] = 'attachment; filename="firmy_export.csv"'
        return response

    def _export_queryset_xlsx(self, queryset):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return HttpResponse('Chyba kniznica openpyxl.', status=500)
        wb = Workbook()
        ws = wb.active
        ws.title = 'Firmy'
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color='DAEEF3', end_color='DAEEF3', fill_type='solid')
        headers = [label for _, label in self.EXPORT_FIELDS]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = header_font
            cell.fill = header_fill
        for row_idx, row_data in enumerate(self._get_export_data(queryset), 2):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=str(value) if value else '')
        for col in ws.columns:
            max_length = max(len(str(cell.value or '')) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_length + 2, 50)
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="firmy_export.xlsx"'
        return response

    def add_company_from_ruz_view(self, request):
        from django.shortcuts import render, redirect
        if request.method == 'POST':
            ico = request.POST.get('ico', '').strip()
            if not ico:
                self.message_user(request, 'ICO je povinne.', messages.ERROR)
                return redirect('admin:companies_company_changelist')
            if Company.objects.filter(ico=ico).exists():
                self.message_user(request, f'Firma s ICO {ico} uz existuje v databaze.', messages.WARNING)
                return redirect('admin:companies_company_changelist')
            try:
                from registers.tasks import sync_single_company_from_ruz
                sync_single_company_from_ruz.delay(ico)
                self.message_user(request, f'Import firmy s ICO {ico} z RUZ bol naplanovany.', messages.SUCCESS)
            except Exception as celery_error:
                try:
                    from registers.tasks import sync_single_company_from_ruz
                    result = sync_single_company_from_ruz(ico)
                    self.message_user(request, f'Import dokonceny: {result}', messages.SUCCESS)
                except Exception as e:
                    self.message_user(request, f'Chyba pri importe: {str(e)}', messages.ERROR)
            return redirect('admin:companies_company_changelist')
        context = {
            **self.admin_site.each_context(request),
            'title': 'Pridat firmu z RUZ API',
            'opts': self.model._meta,
        }
        return render(request, 'admin/companies/add_from_ruz.html', context)

    def trigger_full_ruz_sync_view(self, request):
        from django.shortcuts import redirect
        from registers.tasks import fetch_ruz_data_task
        if request.method == 'POST':
            try:
                fetch_ruz_data_task.delay()
                self.message_user(
                    request,
                    'Kompletna synchronizacia z RUZ API bola spustena.',
                    messages.SUCCESS
                )
            except Exception as e:
                self.message_user(
                    request,
                    f'Celery/Redis nie je dostupny. Chyba: {str(e)[:100]}',
                    messages.ERROR
                )
        return redirect('admin:companies_company_changelist')

    def changelist_view(self, request, extra_context=None):
        from registers.models import SyncProgress

        extra_context = extra_context or {}
        extra_context['show_ruz_buttons'] = True

        stats_cache_key = 'companies_admin_dashboard_stats_v2'
        options_cache_key = 'companies_admin_combined_filter_options_v2'

        dashboard_stats = cache.get(stats_cache_key)
        if dashboard_stats is None:
            debt_q = Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0)
            counts = Company.objects.aggregate(
                total=Count('id'),
                active=Count('id', filter=Q(datum_zrusenia__isnull=True)),
                with_debts=Count('id', filter=debt_q),
            )
            total_companies = counts.get('total') or 0
            active_companies = counts.get('active') or 0
            companies_with_debts = counts.get('with_debts') or 0
            companies_with_orsr = OrsrCompanyProfile.objects.count()
            companies_with_financials = CompanyFinancialResult.objects.values('company_id').distinct().count()
            sync_progress = SyncProgress.objects.filter(sync_type='full').first()

            dashboard_stats = {
                'total': total_companies,
                'active': active_companies,
                'inactive': total_companies - active_companies,
                'with_debts': companies_with_debts,
                'with_orsr': companies_with_orsr,
                'without_orsr': max(total_companies - companies_with_orsr, 0),
                'with_financials': companies_with_financials,
                'without_financials': max(total_companies - companies_with_financials, 0),
                'sync_progress': sync_progress,
                'orsr_pct': round(companies_with_orsr / total_companies * 100, 1) if total_companies else 0,
                'fin_pct': round(companies_with_financials / total_companies * 100, 1) if total_companies else 0,
                'total_display': format_int_space(total_companies),
                'active_display': format_int_space(active_companies),
                'with_debts_display': format_int_space(companies_with_debts),
            }
            cache.set(stats_cache_key, dashboard_stats, 120)

        combined_filter_options = cache.get(options_cache_key)
        if combined_filter_options is None:
            tax_reliability_values = list(
                Company.objects.exclude(tax_reliability__isnull=True)
                .exclude(tax_reliability__exact='')
                .values_list('tax_reliability', flat=True)
                .distinct()
                .order_by('tax_reliability')
            )
            kraj_values = list(
                Company.objects.exclude(kraj__isnull=True)
                .exclude(kraj__exact='')
                .values_list('kraj', flat=True)
                .distinct()
                .order_by('kraj')
            )
            size_values = list(
                Company.objects.exclude(velkost_organizacie__isnull=True)
                .exclude(velkost_organizacie__exact='')
                .values_list('velkost_organizacie', flat=True)
                .distinct()
                .order_by('velkost_organizacie')
            )
            nace_values = list(
                Company.objects.exclude(sk_NACE__isnull=True)
                .exclude(sk_NACE__exact='')
                .values_list('sk_NACE', flat=True)
                .distinct()
                .order_by('sk_NACE')
            )
            founding_years = [
                d.year
                for d in Company.objects.exclude(datum_zalozenia__isnull=True).dates(
                    'datum_zalozenia',
                    'year',
                    order='DESC',
                )
            ]
            combined_filter_options = {
                # Use canonical choices prepared in models.py
                'legal_forms': LEGAL_FORMS_CHOICES,
                'kraje': kraj_values,
                'velkosti': size_values,
                'nace_codes': nace_values,
                'tax_reliability': tax_reliability_values,
                'founded_years': founding_years,
                'vat_payer': [('1', 'Áno'), ('0', 'Nie')],
                'debt_state': [('has_debt', 'S dlhmi'), ('debt_free', 'Bez dlhov')],
                'data_state': [
                    ('complete', 'ORSR + financie'),
                    ('missing_orsr', 'Chýba ORSR'),
                    ('missing_financials', 'Chýbajú financie'),
                ],
            }
            cache.set(options_cache_key, combined_filter_options, 300)

        extra_context['dashboard_stats'] = dashboard_stats
        extra_context['combined_filter_options'] = combined_filter_options

        return super().changelist_view(request, extra_context=extra_context)


@admin.register(SectorBenchmark)
class SectorBenchmarkAdmin(UnfoldModelAdmin):
    list_display = [
        'nace_section', 'year', 'company_count',
        'median_revenue', 'median_roa', 'median_roe',
        'computed_at',
    ]
    list_filter = ['nace_section', 'year']
    ordering = ['-year', 'nace_section']
    readonly_fields = [
        f.name for f in SectorBenchmark._meta.fields
    ]

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False


@admin.register(PostalCodeArea)
class PostalCodeAreaAdmin(UnfoldModelAdmin):
    """Read-only mirror of the MV SR address import.

    The table is derived wholesale from a published dataset, so editing a row
    here would be undone by the next import without anyone noticing -- the same
    reason `SectorBenchmark` is read-only. It is registered at all so the two
    numbers that decide whether a seat gets a map can be inspected: `point_count`
    (how many address points the centroid is measured from) and `radius_m` (how
    far the pin's claim actually reaches).
    """
    list_display = [
        'psc', 'dominant_obec', 'okres', 'kraj',
        'radius_m', 'point_count', 'obec_count',
        'source_version', 'imported_at',
    ]
    list_filter = ['kraj', 'okres']
    search_fields = ['psc', 'dominant_obec']
    ordering = ['psc']
    readonly_fields = [
        f.name for f in PostalCodeArea._meta.fields
    ]

    def has_add_permission(self, request):
        return False

    def has_delete_permission(self, request, obj=None):
        return False

    def has_change_permission(self, request, obj=None):
        return False
