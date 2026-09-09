"""Admin company CRUD with rich filtering."""

from __future__ import annotations

import csv
import hashlib
import json
import logging
import time
import uuid
from io import BytesIO

from django.core.cache import cache
from django.db.models import Avg, Count, Exists, OuterRef, Q, Subquery, Sum
from django.http import HttpResponse, StreamingHttpResponse
from rest_framework import status, viewsets
from rest_framework.decorators import action
from rest_framework.pagination import PageNumberPagination
from rest_framework.response import Response

from adminapi.permissions import IsAdminStaff
from adminapi.serializers import (
    AdminCompanyDetailSerializer,
    AdminCompanyListSerializer,
)
from adminapi.services import CompanyFilterService
from companies.models import Company, CompanyFinancialResult
from registers.models import CompanySyncStatus, OrsrCompanyProfile

logger = logging.getLogger(__name__)


class _CompanyPagination(PageNumberPagination):
    page_size = 50
    page_size_query_param = "page_size"
    max_page_size = 200


class _CSVBuffer:
    """Minimal csv.writer target that returns each generated row for streaming."""

    def write(self, value):
        return value


class AdminCompanyViewSet(viewsets.ModelViewSet):
    permission_classes = [IsAdminStaff]
    queryset = Company.objects.all()
    pagination_class = _CompanyPagination
    filter_service = CompanyFilterService()
    MAX_SYNC_EXPORT_ROWS = 5_000
    REPORT_CACHE_TIMEOUT = 60
    REPORT_CACHE_LOCK_TIMEOUT = 30
    REPORT_CACHE_POLL_INTERVAL = 0.1
    REPORT_CACHE_CONTROL_PARAMS = {
        "export",
        "format",
        "light",
        "mode",
        "page",
        "page_size",
        "preset",
    }

    EXPORT_FIELDS = [
        ("ico", "IČO"),
        ("nazov_UJ", "Názov"),
        ("mesto", "Mesto"),
        ("psc", "PSČ"),
        ("kraj", "Kraj"),
        ("okres", "Okres"),
        ("pravna_forma", "Právna forma"),
        ("sk_NACE", "SK NACE"),
        ("velkost_organizacie", "Veľkosť"),
        ("datum_zalozenia", "Založená"),
        ("datum_zrusenia", "Zrušená"),
        ("vat_payer", "Platiteľ DPH"),
        ("tax_reliability", "Daňová spoľahlivosť"),
        ("tax_debt", "Daňový dlh"),
        ("debt_vszp", "Dlh VSZP"),
        ("debt_soc_poist", "Dlh SP"),
        ("has_orsr", "Má ORSR"),
        ("has_financials", "Má financie"),
        ("debt_state", "Stav dlhov"),
        ("latest_financial_year", "Rok financií"),
        ("latest_revenue", "Tržby"),
        ("latest_profit", "Zisk"),
        ("lead_score", "Lead score"),
        ("lead_confidence", "Confidence"),
        ("sync_failures", "Sync chyby"),
        ("is_blocked", "Blokované"),
    ]

    def get_serializer_class(self):
        if self.action == "list":
            return AdminCompanyListSerializer
        return AdminCompanyDetailSerializer

    def _base_queryset(self):
        return Company.objects.all().select_related("orsr_profile", "lead_score", "enrichment")

    def _listing_queryset(self):
        """Always fully annotate queryset to avoid DISTINCT() calls during filtering."""
        latest_financial = CompanyFinancialResult.objects.filter(company_id=OuterRef("pk")).order_by(
            "-year",
            "-updated_at",
            "-id",
        )
        return self._base_queryset().annotate(
            # Existence checks - these eliminate need for DISTINCT when filtering
            has_orsr_flag=Exists(OrsrCompanyProfile.objects.filter(company_id=OuterRef("pk"))),
            has_financials_flag=Exists(CompanyFinancialResult.objects.filter(company_id=OuterRef("pk"))),
            # Financial data
            latest_financial_year=Subquery(latest_financial.values("year")[:1]),
            latest_revenue=Subquery(latest_financial.values("revenue")[:1]),
            latest_profit=Subquery(latest_financial.values("profit")[:1]),
            # Sync status aggregation
            sync_failures=Subquery(
                CompanySyncStatus.objects.filter(company_id=OuterRef("pk"))
                .order_by("-consecutive_failures", "-updated_at", "-id")
                .values("consecutive_failures")[:1]
            ),
            is_blocked_flag=Exists(CompanySyncStatus.objects.filter(company_id=OuterRef("pk"), is_blocked=True)),
            # Profit/revenue state helpers - computed fields to avoid joins during filtering
            has_profit=Exists(
                CompanyFinancialResult.objects.filter(company_id=OuterRef("pk"), profit__gt=0)
            ),
            has_loss=Exists(
                CompanyFinancialResult.objects.filter(company_id=OuterRef("pk"), profit__lt=0)
            ),
            has_revenue=Exists(
                CompanyFinancialResult.objects.filter(company_id=OuterRef("pk"), revenue__gt=0)
            ),
        )

    def _filtered_queryset(self):
        return self.filter_service.apply(self._listing_queryset(), self.request.query_params)

    def _filter_params(self):
        params = self.request.query_params.copy()
        builder = params.get("filter_builder")
        if builder and not isinstance(builder, str):
            params["filter_builder"] = builder
        return params

    def get_queryset(self):
        if self.action == "retrieve":
            return self._base_queryset().prefetch_related("financial_results", "sync_statuses")
        return self.filter_service.apply(self._listing_queryset(), self._filter_params())

    def _normalized_report_filters(self, params):
        """Return the effective filtering inputs in a canonical form."""
        normalized = {}
        for key in sorted(params.keys()):
            # These controls do not affect the aggregate queryset. The report
            # mode and authenticated requester are represented separately.
            if key in self.REPORT_CACHE_CONTROL_PARAMS:
                continue
            if hasattr(params, "getlist"):
                values = params.getlist(key)
                value = values[-1] if values else ""
            else:
                value = params.get(key, "")
                if isinstance(value, (list, tuple)):
                    value = value[-1] if value else ""
            if key == "filter_builder" and isinstance(value, str):
                try:
                    value = json.dumps(
                        json.loads(value),
                        ensure_ascii=False,
                        separators=(",", ":"),
                        sort_keys=True,
                    )
                except (TypeError, ValueError):
                    # Invalid builders are handled by the filter service; keep
                    # their raw value isolated from valid filter cache entries.
                    pass
            normalized[key] = str(value)
        return normalized

    def _report_cache_key(self, request, mode="full", filters=None):
        """Build an opaque key scoped to the effective filters and staff user."""
        user = getattr(request, "user", None)
        auth_scope = {
            "user_id": (
                str(getattr(user, "pk", ""))
                if getattr(user, "is_authenticated", False)
                else None
            ),
            "is_staff": bool(getattr(user, "is_staff", False)),
            "is_superuser": bool(getattr(user, "is_superuser", False)),
        }
        payload = {
            "auth": auth_scope,
            "filters": self._normalized_report_filters(
                filters if filters is not None else request.query_params
            ),
            "mode": mode,
            "version": 4,
        }
        raw = json.dumps(
            payload,
            ensure_ascii=False,
            separators=(",", ":"),
            sort_keys=True,
        )
        digest = hashlib.sha256(raw.encode("utf-8")).hexdigest()
        return f"admin-company-report:v4:{digest}"

    def _release_report_lock(self, lock_key: str, lock_token: str) -> None:
        """Atomically delete this worker's Redis lock when the backend permits.

        Django's public cache API doesn't expose compare-and-delete. Its built-in
        Redis cache backend does expose the required Redis client internally, so
        use a Lua comparison there. Other cache backends retain the short TTL:
        deleting through the generic API after a separate read could erase a
        newly acquired lock belonging to another worker.
        """
        try:
            backend_client = getattr(cache, "_cache", None)
            get_client = getattr(backend_client, "get_client", None)
            serializer = getattr(backend_client, "_serializer", None)
            make_key = getattr(cache, "make_key", None)
            if not (
                callable(get_client)
                and callable(getattr(serializer, "dumps", None))
                and callable(make_key)
            ):
                return

            redis_key = make_key(lock_key)
            redis_client = get_client(redis_key, write=True)
            compare_and_delete = getattr(redis_client, "eval", None)
            if not callable(compare_and_delete):
                return
            compare_and_delete(
                (
                    "if redis.call('get', KEYS[1]) == ARGV[1] "
                    "then return redis.call('del', KEYS[1]) else return 0 end"
                ),
                1,
                redis_key,
                serializer.dumps(lock_token),
            )
        except Exception:
            # Retaining the short lock TTL is safer than using an unsafe
            # read-then-delete fallback, and must not hide an aggregate error.
            logger.warning(
                "Admin company report generation lock could not be released; "
                "it will expire normally.",
                exc_info=True,
            )

    def _is_light_mode(self, request):
        mode = (request.query_params.get("mode") or "").strip().lower()
        if mode == "full":
            return False
        if mode == "light":
            return True
        flag = (request.query_params.get("light") or "").strip().lower()
        if flag in {"1", "true", "yes"}:
            return True
        # Default to light mode to keep the list endpoint responsive on very large datasets.
        return True

    @action(detail=False, methods=["get"])
    def presets(self, request):
        return Response(self.filter_service.get_presets())

    @action(detail=False, methods=["get"])
    def report(self, request):
        export_format = request.query_params.get("export")
        filter_params = request.query_params.copy()
        preset_key = request.query_params.get("preset")
        if preset_key:
            preset = self.filter_service.get_preset(preset_key)
            if preset:
                filter_params.update(preset.get("filters", {}))
        queryset = self.filter_service.apply(self._listing_queryset(), filter_params)
        if export_format in {"csv", "xlsx"}:
            return self._export_queryset(queryset, export_format)

        is_light_mode = self._is_light_mode(request)
        mode = "light" if is_light_mode else "full"
        cache_key = self._report_cache_key(request, mode=mode, filters=filter_params)
        lock_key = f"{cache_key}:lock"
        lock_token = uuid.uuid4().hex

        try:
            cached_summary = cache.get(cache_key)
            if cached_summary is not None:
                return Response(self._report_response(cached_summary, request))
            owns_generation_lock = cache.add(
                lock_key,
                lock_token,
                timeout=self.REPORT_CACHE_LOCK_TIMEOUT,
            )
        except Exception:
            logger.warning(
                "Admin company report cache is unavailable; computing an uncached report.",
                exc_info=True,
            )
            return Response(
                self._build_aggregate_report(queryset, is_light_mode, request)
            )

        if not owns_generation_lock:
            deadline = time.monotonic() + self.REPORT_CACHE_LOCK_TIMEOUT
            while time.monotonic() < deadline:
                time.sleep(self.REPORT_CACHE_POLL_INTERVAL)
                try:
                    cached_summary = cache.get(cache_key)
                except Exception:
                    logger.warning(
                        "Admin company report cache failed while waiting; computing an uncached report.",
                        exc_info=True,
                    )
                    return Response(
                        self._build_aggregate_report(queryset, is_light_mode, request)
                    )
                if cached_summary is not None:
                    return Response(self._report_response(cached_summary, request))
            return Response(
                {"detail": "Report is already being generated. Please retry shortly."},
                status=status.HTTP_503_SERVICE_UNAVAILABLE,
            )

        try:
            summary = self._build_aggregate_report(queryset, is_light_mode, request)
            try:
                cache.set(cache_key, summary, self.REPORT_CACHE_TIMEOUT)
            except Exception:
                logger.warning(
                    "Admin company report was generated but could not be cached.",
                    exc_info=True,
                )
            return Response(self._report_response(summary, request))
        finally:
            self._release_report_lock(lock_key, lock_token)

    def _report_response(self, summary, request):
        """Restore request-specific metadata omitted from the shared cache."""
        return {
            **summary,
            "filters_applied": dict(request.query_params),
        }

    def _build_aggregate_report(self, queryset, is_light_mode, request):
        debt_q = Q(debt_vszp__gt=0) | Q(debt_soc_poist__gt=0) | Q(tax_debt__gt=0)
        counts = queryset.aggregate(
            count=Count("id"),
            active_count=Count("id", filter=Q(datum_zrusenia__isnull=True)),
            inactive_count=Count("id", filter=Q(datum_zrusenia__isnull=False)),
            debt_free_count=Count("id", filter=~debt_q),
        )

        if not is_light_mode:
            counts["has_orsr_count"] = queryset.filter(has_orsr_flag=True).count()
            counts["has_financials_count"] = queryset.filter(has_financials_flag=True).count()
            counts["blocked_count"] = queryset.filter(is_blocked_flag=True).count()
        summary = {
            "count": counts.get("count") or 0,
            "active_count": counts.get("active_count") or 0,
            "inactive_count": counts.get("inactive_count") or 0,
            "debt_free_count": counts.get("debt_free_count") or 0,
            "has_orsr_count": counts.get("has_orsr_count") or 0,
            "has_financials_count": counts.get("has_financials_count") or 0,
            "blocked_count": counts.get("blocked_count") or 0,
            "avg_lead_score": 0,
            "avg_confidence": 0,
            "revenue_sum": 0,
            "profit_sum": 0,
            "tax_debt_sum": 0,
            "debt_vszp_sum": 0,
            "debt_soc_poist_sum": 0,
            "lead_score_sum": 0,
            "report_mode": "light" if is_light_mode else "full",
            "presets": self.filter_service.get_presets(),
            "top_companies": [],
        }

        if not is_light_mode:
            score_aggs = queryset.aggregate(
                avg_lead_score=Avg("lead_score__score"),
                avg_confidence=Avg("enrichment__confidence_score"),
                lead_score_sum=Sum("lead_score__score"),
            )
            debt_sums = queryset.aggregate(
                tax_debt_sum=Sum("tax_debt"),
                debt_vszp_sum=Sum("debt_vszp"),
                debt_soc_poist_sum=Sum("debt_soc_poist"),
            )
            financial_sums = queryset.aggregate(
                revenue_sum=Sum("latest_revenue"),
                profit_sum=Sum("latest_profit"),
            )
            summary.update({
                "avg_lead_score": round(score_aggs.get("avg_lead_score") or 0, 2),
                "avg_confidence": round(score_aggs.get("avg_confidence") or 0, 2),
                "revenue_sum": financial_sums.get("revenue_sum") or 0,
                "profit_sum": financial_sums.get("profit_sum") or 0,
                "tax_debt_sum": debt_sums.get("tax_debt_sum") or 0,
                "debt_vszp_sum": debt_sums.get("debt_vszp_sum") or 0,
                "debt_soc_poist_sum": debt_sums.get("debt_soc_poist_sum") or 0,
                "lead_score_sum": score_aggs.get("lead_score_sum") or 0,
                "top_companies": AdminCompanyListSerializer(
                    queryset.order_by("-lead_score__score", "-id")[:10],
                    many=True,
                ).data,
            })
        return summary

    def _export_queryset(self, queryset, export_format: str):
        export_count = queryset.count()
        if export_count > self.MAX_SYNC_EXPORT_ROWS:
            return Response(
                {
                    "detail": (
                        f"Export je obmedzený na {self.MAX_SYNC_EXPORT_ROWS} riadkov. "
                        "Zúžte filtre a skúste znova."
                    ),
                    "max_rows": self.MAX_SYNC_EXPORT_ROWS,
                },
                status=status.HTTP_400_BAD_REQUEST,
            )

        export_queryset = queryset.order_by("-id")[:self.MAX_SYNC_EXPORT_ROWS]
        if export_format == "csv":
            return self._export_csv(export_queryset)
        rows = AdminCompanyListSerializer(export_queryset, many=True).data
        return self._export_xlsx(rows)

    def _export_csv(self, queryset):
        def stream_rows():
            writer = csv.writer(_CSVBuffer(), delimiter=";")
            yield "\ufeff"
            yield writer.writerow([label for _, label in self.EXPORT_FIELDS])
            for company in queryset.iterator(chunk_size=500):
                row = AdminCompanyListSerializer(company).data
                yield writer.writerow([
                    self._export_value(row.get(field))
                    for field, _ in self.EXPORT_FIELDS
                ])

        response = StreamingHttpResponse(stream_rows(), content_type="text/csv; charset=utf-8")
        response["Content-Disposition"] = 'attachment; filename="companies-report.csv"'
        return response

    def _export_xlsx(self, rows):
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill
        except ImportError:
            return Response({"detail": "openpyxl is required for XLSX export."}, status=status.HTTP_500_INTERNAL_SERVER_ERROR)

        wb = Workbook()
        ws = wb.active
        ws.title = "Companies"
        header_font = Font(bold=True)
        header_fill = PatternFill(start_color="DAEEF3", end_color="DAEEF3", fill_type="solid")
        for col, (_, label) in enumerate(self.EXPORT_FIELDS, 1):
            cell = ws.cell(row=1, column=col, value=label)
            cell.font = header_font
            cell.fill = header_fill

        for row_idx, row in enumerate(rows, 2):
            for col_idx, (field, _) in enumerate(self.EXPORT_FIELDS, 1):
                ws.cell(row=row_idx, column=col_idx, value=self._export_value(row.get(field)))

        for column in ws.columns:
            max_length = max(len(str(cell.value or "")) for cell in column)
            ws.column_dimensions[column[0].column_letter].width = min(max_length + 2, 50)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        response = HttpResponse(
            output.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="companies-report.xlsx"'
        return response

    def _export_value(self, value):
        if value is None:
            return ""
        if isinstance(value, bool):
            return "Áno" if value else "Nie"
        return value
